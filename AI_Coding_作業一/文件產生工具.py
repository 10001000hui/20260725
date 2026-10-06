"""重新產生作業 Word 文件。需要 python-docx；不會讀取私人部署資訊。"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent

def configure(doc):
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[name]
        style.font.name = 'Microsoft JhengHei'
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('AD462E')
    footer = section.footer.paragraphs[0]
    footer.text = '築跡 Studio Notes｜AI Coding 作業一｜2026-10-06'
    footer.style = 'Caption'

def plain(text):
    return text.replace('**', '').replace('`', '')

def convert(path):
    doc = Document()
    configure(doc)
    code = False
    for line in path.read_text().splitlines():
        if line.startswith('```'):
            code = not code
            continue
        if not line.strip():
            continue
        if code:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
        elif line.startswith('# '):
            doc.add_heading(plain(line[2:]), 0)
        elif line.startswith('## '):
            doc.add_heading(plain(line[3:]), 1)
        elif line.startswith('### '):
            doc.add_heading(plain(line[4:]), 2)
        elif line.startswith('- '):
            doc.add_paragraph(plain(line[2:]), style='List Bullet')
        elif re.match(r'^\d+\. ', line):
            doc.add_paragraph(plain(re.sub(r'^\d+\. ', '', line)), style='List Number')
        elif line.startswith('|'):
            if '---' not in line:
                doc.add_paragraph(' ／ '.join(plain(x.strip()) for x in line.strip('|').split('|')))
        else:
            doc.add_paragraph(plain(line))
    if path.name.startswith('01_'):
        doc.add_page_break()
        doc.add_heading('網站畫面（本地正式模式）', 1)
        doc.add_paragraph('此圖取自本地 Chromium，供了解已完成的版面；不代表公開網址驗證通過。完整桌面、平板與手機畫面另存於驗證資料。')
        doc.add_picture(str(ROOT / '驗證資料' / '首頁畫面.png'), width=Inches(6.3))
    doc.save(path.with_suffix('.docx'))

if __name__ == '__main__':
    for name in ['01_一步一步操作說明.md', '02_作業繳交內容.md', '03_驗證報告.md']:
        convert(ROOT / name)
        print('已產生：' + Path(name).with_suffix('.docx').name)
